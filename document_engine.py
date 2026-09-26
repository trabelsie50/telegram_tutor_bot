"""
document_engine.py - استخراج النص من ملفات PDF و Word (DOCX) لإضافتها إلى قاعدة RAG.

يدعم:
- PDF نصي (native) عبر pypdf.
- PDF ممسوح (scanned / صور) عبر OCR تلقائي كحل احتياطي (pdf2image + ocr_engine).
- DOCX (فقرات وجداول) عبر python-docx. لا يدعم .doc القديم.
"""

import logging
import os
from typing import Any, Dict, List

from pypdf import PdfReader
import docx  # python-docx

from ocr_engine import extract_text_from_image

logger = logging.getLogger(__name__)

# تحت هذا الحد من عدد الحروف نعتبر صفحة PDF ممسوحة (بلا نص قابل للاستخراج) وتحتاج OCR
MIN_CHARS_PER_PAGE_NATIVE = 20


def extract_text_from_pdf(pdf_path: str) -> Dict[str, Any]:
    """يستخرج النص من ملف PDF: نص مباشر إن وُجد، وإلا OCR للصفحات الممسوحة.

    يُرجع قاموساً يحتوي على: text, page_count, ocr_pages_count, needs_review.
    """
    if not os.path.isfile(pdf_path):
        raise FileNotFoundError(f"ملف PDF غير موجود: {pdf_path}")

    reader = PdfReader(pdf_path)
    pages_text: List[str] = []
    scanned_pages: List[int] = []

    for i, page in enumerate(reader.pages):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            logger.warning("فشل استخراج نص الصفحة %d من PDF: %s", i, exc)
            text = ""

        if len(text) < MIN_CHARS_PER_PAGE_NATIVE:
            scanned_pages.append(i)
        pages_text.append(text)

    needs_ocr = len(scanned_pages) > 0

    if needs_ocr:
        try:
            from pdf2image import convert_from_path

            images = convert_from_path(pdf_path)
            for i in scanned_pages:
                if i >= len(images):
                    continue
                tmp_img_path = f"{pdf_path}_page_{i}.jpg"
                images[i].save(tmp_img_path, "JPEG")
                try:
                    ocr_result = extract_text_from_image(tmp_img_path)
                    pages_text[i] = ocr_result.get("text", "")
                finally:
                    if os.path.exists(tmp_img_path):
                        os.remove(tmp_img_path)
        except Exception as exc:
            logger.error("فشل معالجة صفحات PDF الممسوحة بـ OCR (poppler مثبّت؟): %s", exc)

    full_text = "\n\n".join(t for t in pages_text if t.strip())

    return {
        "text": full_text,
        "page_count": len(reader.pages),
        "ocr_pages_count": len(scanned_pages),
        "needs_review": len(scanned_pages) > 0,
    }


def extract_text_from_docx(docx_path: str) -> Dict[str, Any]:
    """يستخرج النص (فقرات + جداول) من ملف Word بصيغة .docx."""
    if not os.path.isfile(docx_path):
        raise FileNotFoundError(f"ملف Word غير موجود: {docx_path}")

    document = docx.Document(docx_path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    table_lines: List[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))

    full_text = "\n".join(paragraphs + table_lines)

    return {
        "text": full_text,
        "paragraph_count": len(paragraphs),
        "needs_review": len(full_text.strip()) < 20,
    }


def extract_text_from_document(file_path: str) -> Dict[str, Any]:
    """نقطة دخول موحّدة: تحدد نوع الملف من الامتداد وتستخرج النص المناسب.

    يرفع ValueError إذا كانت الصيغة غير مدعومة (.doc القديم مثلاً).
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(
            f"صيغة ملف غير مدعومة: {ext}. المدعوم حالياً: PDF و Word (.docx فقط)."
        )
